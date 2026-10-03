"""Build docs/pitch/AirPower-SIH2026-Idea-Presentation.pptx on the official SIH idea-presentation format.

Keeps the template's slides, headings, colours, header/footer and logo; removes the instructions
slide (the template says to delete it); fills the six slides with AirPower content, app screenshots
and diagrams. Benchmark figures come from benchmarks/results/results.csv (CLAUDE.md rule 3).

    python docs/pitch/build_deck.py "<SIH idea format>.pptx" <screenshots dir>
"""

from __future__ import annotations

import csv
import io
import statistics as st
import sys
from pathlib import Path

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "AirPower-SIH2026-Idea-Presentation.pptx"
IMG = Path(sys.argv[2])
REPO = "https://github.com/Suyash2527/airpower-ops-optimiser"
TESTS = 385  # backend pytest result at build time (pytest -q: 385 passed)

# ------------------------------------------------------------------ benchmark numbers (from the CSV)
rows = list(csv.DictReader(open(ROOT / "benchmarks" / "results" / "results.csv", encoding="utf-8")))
by: dict[str, list[dict]] = {}
for r in rows:
    by.setdefault(r["planner"], []).append(r)
pwc = {k: [float(r["priority_weighted_coverage"]) for r in v] for k, v in by.items()}
N = len(pwc["cpsat"])
gain_g = st.mean((c - g) / g for c, g in zip(pwc["cpsat"], pwc["greedy"], strict=True) if g > 0)
gain_f = st.mean((c - f) / f for c, f in zip(pwc["cpsat"], pwc["fifo"], strict=True) if f > 0)
better_g = sum(c > g + 1e-9 for c, g in zip(pwc["cpsat"], pwc["greedy"], strict=True))
worse_g = sum(c < g - 1e-9 for c, g in zip(pwc["cpsat"], pwc["greedy"], strict=True))
n_opt = sum(r["solver_status"] == "OPTIMAL" for r in by["cpsat"])
viol = sum(int(r["violations"]) for r in rows)
PLANS = len(rows)
print(f"benchmarks: N={N} greedy +{gain_g:.1%} fifo +{gain_f:.1%} better {better_g} worse {worse_g} optimal {n_opt} viol {viol}/{PLANS}")

# ------------------------------------------------------------------ style
BLUE, GREEN, RED, SKY, ORANGE = "1F497D", "2E8B57", "B3372F", "0070C0", "E98B1D"
INK, MUTED, TINT, EDGE = "000000", "595959", "F2F6FA", "D5DEE9"
FONT = "Arial"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def para(tf, first, runs, size=11, align=None, after=2, bullet=None, indent=0.16, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align:
        p.alignment = align
    p.space_after = Pt(after)
    p.space_before = Pt(0)
    if line:
        p.line_spacing = line
    if bullet:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": SKY}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": bullet}))
    for item in runs:
        text, o = (item, {}) if isinstance(item, str) else item
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", False)
        f.italic = o.get("italic", False)
        f.underline = o.get("underline", False)
        f.color.rgb = rgb(o.get("color", INK))
    return p


def tbox(slide, x, y, w, h, paras, size=11, anchor=MSO_ANCHOR.TOP, margin=0.04, **kw):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, pr in enumerate(paras):
        para(tf, i == 0, pr, size=size, **kw)
    return s


def box(slide, x, y, w, h, fill=TINT, line=EDGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label_box(slide, x, y, w, h, text, fill, size=11, color="FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = box(slide, x, y, w, h, fill=fill, line=None, shape=shape)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, True, [(text, {"bold": True, "color": color, "size": size})], align=PP_ALIGN.CENTER, after=0)
    return s


def heading(slide, x, y, w, text, major=True):
    mark = "❖ " if major else "• "
    return tbox(slide, x, y, w, 0.34,
                [[(mark + text, {"bold": True, "color": BLUE, "size": 13 if major else 12, "underline": major})]], after=0)


def pic(slide, name, x, y, w, border=True):
    path = IMG / f"{name}.png"
    iw, ih = Image.open(path).size
    h = w * ih / iw
    p = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        p.line.color.rgb = rgb(EDGE)
        p.line.width = Pt(1)
    return p, h


def qr(slide, url, x, y, size):
    img = qrcode.make(url, border=1, box_size=12)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(size), Inches(size))


def drop(slide, name):
    for s in list(slide.shapes):
        if s.name == name:
            s._element.getparent().remove(s._element)


def set_title(slide, text, size=None):
    t = slide.shapes.title
    r_all = t._element.txBody.findall(".//" + qn("a:r"))
    r_all[0].find(qn("a:t")).text = text
    for r in r_all[1:]:
        r.getparent().remove(r)
    for br in t._element.txBody.findall(".//" + qn("a:br")):
        br.getparent().remove(br)
    if size:
        r_all[0].find(qn("a:rPr")).set("sz", str(int(size * 100)))


def oval_team(slide, text="CodeVerse"):
    for s in slide.shapes:
        if s.name.startswith("Oval"):
            s.text_frame.paragraphs[0].runs[0].text = text
            s.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
            s.text_frame.word_wrap = False


# ------------------------------------------------------------------ build
prs = Presentation(sys.argv[1])
s1, s2, s3, s4, s5, s6, s7 = prs.slides

# the template says to delete the instructions slide before upload
sld_ids = prs.slides._sldIdLst
last = sld_ids[-1]
prs.part.drop_rel(last.rId)
sld_ids.remove(last)

# ---- 1 · Title page
for r in s1.shapes.title._element.txBody.findall(".//" + qn("a:r")):
    if r.find(qn("a:t")).text.strip() == "2025":
        r.find(qn("a:t")).text = "2026"
box1 = next(s for s in s1.shapes if s.name == "TextBox 9")
txb = box1._element.txBody
for p in txb.findall(qn("a:p")):
    txb.remove(p)
txb.append(txb.makeelement(qn("a:p"), {}))
items = [
    ("Problem Statement ID – ", "SIH26250"),
    ("Problem Statement Title- ", "Planning and dynamically retasking air operations in a contested, rapidly changing environment"),
    ("Theme- ", "Transportation & Logistics"),
    ("PS Category- ", "Software"),
    ("Team ID- ", "133847"),
    ("Team Name (Registered on portal)- ", "CodeVerse"),
    ("GitHub- ", "Suyash2527/airpower-ops-optimiser"),
]
tf = box1.text_frame
for i, (lab, val) in enumerate(items):
    para(tf, i == 0, [(lab, {"bold": True, "size": 17}), (val, {"bold": True, "size": 17, "color": BLUE})],
         bullet="•", indent=0.3, after=5, line=1.1)
box1.top = Inches(2.1)
box1.width = Inches(6.7)

for s in (s2, s3, s4, s5, s6):
    oval_team(s)

# ---- 2 · Idea title / proposed solution
set_title(s2, "AIRPOWER: AI Decision Support for Air Operations", size=24)
s2.shapes.title.left, s2.shapes.title.width = Inches(1.85), Inches(8.8)
drop(s2, "TextBox 8")
heading(s2, 0.35, 1.28, 12.6, "Proposed Solution (Describe your Idea/Solution/Prototype)")
tbox(s2, 0.35, 1.66, 6.3, 0.8, [[
    ("AirPower ", {"bold": True, "color": BLUE}),
    "fuses scattered air-operations data into one live picture, builds an optimised plan with a reason for every "
    "assignment, and proposes ranked re-plans the moment anything changes. ",
    ("A human approves every change.", {"bold": True}),
]], size=11.5, after=0)

heading(s2, 0.35, 2.5, 6.3, "Detailed explanation of the proposed solution", major=False)
steps = [("1", "Fuse", "sources, age, conflicts", BLUE), ("2", "Check", "every constraint per sortie", BLUE),
         ("3", "Optimise", "CP-SAT + baselines", BLUE), ("4", "Approve", "human decision, audited", ORANGE),
         ("5", "Retask", "ranked options + diff", GREEN)]
for i, (n, name, sub, col) in enumerate(steps):
    x = 0.35 + i * 1.27
    box(s2, x, 2.88, 1.19, 1.1, fill=TINT, line=EDGE)
    label_box(s2, x + 0.08, 2.96, 0.3, 0.3, n, col, size=11, shape=MSO_SHAPE.OVAL)
    tbox(s2, x + 0.04, 3.28, 1.12, 0.7, [[(name, {"bold": True, "size": 11.5, "color": col})], [(sub, {"size": 9.5, "color": MUTED})]], after=1)

heading(s2, 0.35, 4.1, 6.3, "How it addresses the problem", major=False)
pairs = [("Data in silos", "One fused picture with source and age"), ("Slow, sub-optimal allocation", "Optimised draft plan, every choice explained"),
         ("Painful re-planning", "Ranked options with a plan diff")]
for i, (a, b) in enumerate(pairs):
    y = 4.5 + i * 0.4
    label_box(s2, 0.35, y, 2.2, 0.32, a, RED, size=10)
    tbox(s2, 2.58, y, 0.3, 0.32, [[("→", {"bold": True, "size": 14, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE, after=0)
    label_box(s2, 2.88, y, 3.77, 0.32, b, GREEN, size=10)

heading(s2, 0.35, 5.75, 6.3, "Innovation and uniqueness of the solution", major=False)
tbox(s2, 0.35, 6.07, 6.4, 0.85, [
    [("Explainable by construction: reason code + sentence + next-best option", {"size": 10})],
    [("Stability-aware retasking: airborne sorties frozen, minimal-change variants", {"size": 10})],
    [(f"Independent validator: {viol} violations across {PLANS} benchmark plans", {"size": 10})],
], bullet="•", after=1)

p_, h_ = pic(s2, "option", 6.85, 1.7, 6.1)
tbox(s2, 6.85, 1.7 + h_ + 0.03, 6.1, 0.25, [[("Retasking option as the planner sees it: trade-offs and a sortie-level diff (prototype screenshot)", {"size": 9, "color": MUTED, "italic": True})]], after=0)
stats = [(f"+{gain_g:.1%}", f"priority-weighted coverage vs the greedy baseline (mean of {N} scenarios)", SKY),
         (f"+{gain_f:.1%}", "vs first-in first-out planning", GREEN),
         (str(viol), f"validator violations in {PLANS} plans; {TESTS} tests pass", BLUE)]
for i, (big, small, col) in enumerate(stats):
    x = 6.85 + i * 2.05
    box(s2, x, 5.3, 1.95, 1.25, fill="FFFFFF", line=EDGE)
    tbox(s2, x + 0.05, 5.35, 1.85, 0.5, [[(big, {"bold": True, "size": 24, "color": col})]], after=0)
    tbox(s2, x + 0.05, 5.85, 1.85, 0.7, [[(small, {"size": 9, "color": MUTED})]], after=0)
tbox(s2, 6.85, 6.6, 6.1, 0.2, [[("Measured by benchmarks/run_benchmarks.py on synthetic scenarios; see benchmarks/RESULTS.md", {"size": 7.5, "color": MUTED})]], after=0)

# ---- 3 · Technical approach
drop(s3, "TextBox 8")
heading(s3, 0.35, 1.28, 12.6, "Technologies to be used (e.g. programming languages, frameworks, hardware)")
cards = [
    ("Frontend", SKY, "Next.js · React · TypeScript · Tailwind. India map drawn in SVG from open data (no tile server)."),
    ("API and data", BLUE, "Python · FastAPI · Pydantic · SQLModel. SQLite for development, PostgreSQL when deployed."),
    ("Optimisation", GREEN, "Google OR-Tools CP-SAT, with greedy and FIFO baselines; shapely + pyproj for routes and zones."),
    ("Assurance and hardware", ORANGE, f"pytest ({TESTS} tests) and an independent plan validator. No special hardware: a standard CPU."),
]
for i, (name, col, body) in enumerate(cards):
    x = 0.35 + (i % 2) * 4.2
    y = 1.7 + (i // 2) * 1.2
    box(s3, x, y, 4.0, 1.1, fill=TINT, line=EDGE)
    label_box(s3, x + 0.1, y + 0.1, 2.1, 0.28, name, col, size=10)
    tbox(s3, x + 0.08, y + 0.44, 3.85, 0.66, [[(body, {"size": 10})]], after=0)
p_, h_ = pic(s3, "timeline", 8.8, 1.7, 4.15)
tbox(s3, 8.8, 1.7 + h_ + 0.02, 4.15, 0.25, [[("Plan timeline: sorties by aircraft, locked = airborne", {"size": 9, "color": MUTED, "italic": True})]], after=0)

heading(s3, 0.35, 4.2, 12.6, "Methodology and process for implementation (Flow Charts/Images/ working prototype)")
flow = [("1 Ingest", BLUE, "Adapter per source, provenance on every record"),
        ("2 Fuse", BLUE, "Scored on trust, confidence and freshness; conflicts kept, stale flagged"),
        ("3 Check", BLUE, "Range, stock, crew, duty, airspace, weather, threat"),
        ("4 Optimise", BLUE, "CP-SAT within 20 s; greedy plan as warm start"),
        ("5 Approve", ORANGE, "Draft until a person approves; audit entry written"),
        ("6 Retask", GREEN, "Event → affected sorties → ranked options and diff")]
for i, (name, col, body) in enumerate(flow):
    x = 0.35 + i * 2.1
    label_box(s3, x, 4.62, 2.0, 0.42, name, col, size=12, shape=MSO_SHAPE.PENTAGON)
    tbox(s3, x, 5.1, 2.0, 0.85, [[(body, {"size": 10})]], after=0)
box(s3, 0.35, 6.05, 12.6, 0.78, fill="E8F0F8", line=None)
rules = [("Advisory only", "no auto-execution, no targeting logic"), ("Every decision explained", "reason code + plain sentence"), ("Seeded and reproducible", "same seed, same scenario, same result")]
for i, (a, b) in enumerate(rules):
    tbox(s3, 0.5 + i * 4.15, 6.1, 4.0, 0.7, [[(a, {"bold": True, "size": 11, "color": BLUE})], [(b, {"size": 10, "color": INK})]], after=0, anchor=MSO_ANCHOR.MIDDLE)

# ---- 4 · Feasibility and viability
drop(s4, "TextBox 8")
cols = [("Analysis of the feasibility of the idea", GREEN, [
            ("Technical: ", "working end-to-end prototype and test suite"),
            ("Compute: ", "laptop CPU; plans bounded at 20 s; no GPU"),
            ("Data: ", "adapter layer; real feeds replace synthetic ones by configuration"),
            ("Operational: ", "advisory, so command authority is unchanged"),
            ("Security: ", "runs offline and air-gapped; open-source stack")]),
        ("Potential challenges and risks", RED, [
            ("Data access: ", "operational feeds are classified and fragmented"),
            ("Data quality: ", "late or conflicting reports can mislead a plan"),
            ("Trust: ", "planners will not act on a black box"),
            ("Model fidelity: ", "performance and risk figures are placeholders today"),
            ("Adoption: ", "doctrine, training, change management")]),
        ("Strategies for overcoming these challenges", BLUE, [
            ("Accreditation first: ", "on-premise, read-only adapters, role-based access"),
            ("Freshness rules: ", "stale and conflict flags; a human can pin a value"),
            ("Explain and approve: ", "reason codes, plan diff, audit trail"),
            ("Real data: ", "replace placeholders with service-approved figures"),
            ("Shadow mode: ", "run beside today's process before any reliance")])]
for i, (head, col, pts) in enumerate(cols):
    x = 0.35 + i * 4.25
    label_box(s4, x, 1.3, 4.1, 0.42, head, col, size=11.5, shape=MSO_SHAPE.RECTANGLE)
    box(s4, x, 1.72, 4.1, 2.75, fill=TINT, line=None, shape=MSO_SHAPE.RECTANGLE)
    tbox(s4, x + 0.08, 1.78, 3.95, 2.65, [[(a, {"bold": True, "color": col, "size": 11.5}), (b, {"size": 11.5})] for a, b in pts], bullet="•", after=6)

tbox(s4, 0.35, 4.6, 6.2, 0.3, [[("Measured on the prototype (benchmark, not claimed)", {"bold": True, "color": BLUE, "size": 12})]], after=0)
tiles = [(str(N), "seeded scenarios, 3 planners each"), (f"{better_g} of {N}", "scenarios where CP-SAT beat greedy; worse in " + str(worse_g)),
         (f"{n_opt} of {N}", "plans proven optimal within the 20 s limit"), (str(viol), f"violations in {PLANS} validated plans")]
for i, (big, small) in enumerate(tiles):
    x = 0.35 + i * 1.57
    box(s4, x, 4.95, 1.5, 1.4, fill="FFFFFF", line=EDGE)
    tbox(s4, x + 0.04, 5.0, 1.42, 0.5, [[(big, {"bold": True, "size": 20, "color": SKY})]], after=0)
    tbox(s4, x + 0.04, 5.5, 1.42, 0.85, [[(small, {"size": 9, "color": MUTED})]], after=0)
tbox(s4, 0.35, 6.42, 6.2, 0.45, [[("Where the optimiser times out it returns the best plan found; a fallback is always labelled as such. Source: benchmarks/RESULTS.md", {"size": 8.5, "color": MUTED})]], after=0)

tbox(s4, 6.75, 4.6, 6.2, 0.3, [[("Phased rollout and validation plan", {"bold": True, "color": BLUE, "size": 12})]], after=0)
phases = [("Pilot · 0-3 mo", SKY, "1 planning cell, shadow mode, historical data"), ("Trial · 4-9 mo", BLUE, "live read-only feeds, 2 cells, KPIs tracked"),
          ("Expand · 10-18 mo", GREEN, "command-wide; ML models on real data"), ("Scale · Yr 2+", ORANGE, "joint and HADR operations")]
for i, (name, col, body) in enumerate(phases):
    x = 6.75 + i * 1.57
    label_box(s4, x, 4.95, 1.5, 0.4, name, col, size=9.5, shape=MSO_SHAPE.PENTAGON)
    tbox(s4, x, 5.4, 1.5, 0.95, [[(body, {"size": 9.5})]], after=0)
tbox(s4, 6.75, 6.42, 6.2, 0.45, [[("Pilot KPIs: time to first approved plan, re-plan latency, priority-weighted coverage, share of proposals accepted", {"size": 9.5, "color": INK})]], after=0)

# ---- 5 · Impact and benefits
drop(s5, "TextBox 8")
heading(s5, 0.35, 1.28, 7.7, "Potential impact on the target audience")
aud = [("P", "Air-ops planners", "One fused picture and an explained draft plan instead of manual collation", SKY),
       ("C", "Duty controllers", "A disruption becomes affected sorties and ranked options, not a scramble", BLUE),
       ("K", "Commanders", "Coverage by priority, risk and a full audit trail at a glance", GREEN),
       ("H", "HADR / NDMA relief", "Relief airlift allocated by priority and weather with the same engine", ORANGE),
       ("M", "Crew and maintenance", "Duty, rest and maintenance limits enforced at plan time, with the reason", RED)]
for i, (ch, who, what, col) in enumerate(aud):
    y = 1.7 + i * 0.5
    label_box(s5, 0.35, y + 0.03, 0.38, 0.38, ch, col, size=11, shape=MSO_SHAPE.OVAL)
    tbox(s5, 0.82, y, 7.3, 0.46, [[(who + " ", {"bold": True, "size": 11, "color": col}), (what, {"size": 10.5})]], anchor=MSO_ANCHOR.MIDDLE, after=0)

heading(s5, 0.35, 4.3, 7.7, "Benefits of the solution (social, economic, environmental, etc.)")
bens = [("Social", GREEN, ["Faster relief airlift in floods and cyclones", "Less decision stress on planners"]),
        ("Economic", ORANGE, ["Fewer wasted sorties and flying hours", "Open-source stack, no licence cost"]),
        ("Environmental", SKY, ["Less unnecessary flying, so less fuel burned", "Targeted re-plans, not full re-flights"]),
        ("Governance", BLUE, ["Human approval on every change", "Append-only audit trail"])]
for i, (name, col, pts) in enumerate(bens):
    x = 0.35 + i * 1.97
    label_box(s5, x, 4.7, 1.88, 0.34, name, col, size=11, shape=MSO_SHAPE.RECTANGLE)
    box(s5, x, 5.04, 1.88, 1.8, fill=TINT, line=None, shape=MSO_SHAPE.RECTANGLE)
    tbox(s5, x + 0.05, 5.1, 1.78, 1.7, [[(t_, {"size": 11})] for t_ in pts], bullet="•", after=8)

p_, h_ = pic(s5, "map", 8.35, 1.7, 4.6)
tbox(s5, 8.35, 1.7 + h_ + 0.02, 4.6, 0.3, [[("Operating picture: bases, mission areas by priority, threat zones, airspace, sorties. Outline: Natural Earth, not authoritative.", {"size": 8.5, "color": MUTED, "italic": True})]], after=0)
box(s5, 8.35, 5.4, 4.6, 1.42, fill=TINT, line=None, shape=MSO_SHAPE.RECTANGLE)
kpi = tbox(s5, 8.45, 5.45, 4.4, 1.35, [
    [("Pilot KPIs (baseline measured in shadow mode)", {"bold": True, "size": 11, "color": BLUE})],
    [("Time to first approved plan vs the current process", {"size": 10})],
    [("Re-plan latency after a disruption", {"size": 10})],
    [("Priority-weighted coverage; sorties lost to late changes", {"size": 10})],
], after=2)
for p in kpi.text_frame.paragraphs[1:]:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.16))))
    pPr.set("indent", str(-int(Inches(0.16))))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))

# ---- 6 · Research and references
drop(s6, "TextBox 8")
heading(s6, 0.35, 1.28, 12.6, "Details / Links of the reference and research work")
refs = [
    ("Problem and doctrine", [
        "SIH 2026, Problem Statement SIH26250 (Ministry of Defence) · sih.gov.in",
        "US Joint Chiefs of Staff, JP 3-30 Joint Air Operations (joint air tasking cycle) · jcs.mil",
        "NDMA, National Disaster Management Plan: air support for HADR · ndma.gov.in"]),
    ("Human factors and automation", [
        "Endsley (1995). Toward a theory of situation awareness in dynamic systems. Human Factors 37(1) · doi.org/10.1518/001872095779049543",
        "Parasuraman, Sheridan, Wickens (2000). A model for types and levels of human interaction with automation. IEEE Trans. SMC-A 30(3) · doi.org/10.1109/3468.844354"]),
    ("Optimisation and open data", [
        "Google OR-Tools, CP-SAT solver (Apache-2.0) · developers.google.com/optimization",
        "Gopalan and Talluri (1998). Mathematical models in airline schedule planning: a survey. Annals of Operations Research 76",
        "Barnhart et al. (2003). Airline crew scheduling. Handbook of Transportation Science",
        "Natural Earth admin-0, India point of view · OurAirports · Open-Meteo elevation API (open data)"]),
]
for y, (head, items) in zip((1.68, 2.78, 3.98), refs, strict=True):
    tbox(s6, 0.35, y, 8.5, 0.28, [[(head, {"bold": True, "size": 11.5, "color": BLUE})]], after=0)
    tbox(s6, 0.35, y + 0.3, 8.5, 1.2, [[(it, {"size": 10})] for it in items], bullet="•", after=2)
box(s6, 9.15, 1.7, 3.8, 2.95, fill="FFFFFF", line=EDGE)
tbox(s6, 9.15, 1.76, 3.8, 0.3, [[("SEE THE CODE · CHECK THE NUMBERS", {"bold": True, "size": 10.5, "color": BLUE})]], align=PP_ALIGN.CENTER, after=0)
qr(s6, REPO, 10.05, 2.12, 2.0)
tbox(s6, 9.2, 4.14, 3.7, 0.5, [[("Source code", {"bold": True, "size": 10})], [("github.com/Suyash2527/airpower-ops-optimiser", {"size": 8.5, "color": SKY})]], align=PP_ALIGN.CENTER, after=0)
tbox(s6, 0.35, 5.42, 12.6, 0.28, [[("Data used in the prototype (transparency)", {"bold": True, "size": 12, "color": BLUE})]], after=0)
box(s6, 0.35, 5.72, 12.6, 1.12, fill=TINT, line=None, shape=MSO_SHAPE.RECTANGLE)
tbox(s6, 0.45, 5.76, 12.4, 1.06, [
    [("Real open data: ", {"bold": True, "color": BLUE}), "map outline (Natural Earth), alternate civil airfields (OurAirports), base elevations (Open-Meteo)."],
    [("Synthetic data: ", {"bold": True, "color": ORANGE}), "all bases, units, aircraft, crews, missions, threats, weather and events are generated from a seed and labelled SYNTHETIC in the app; performance figures are illustrative."],
    [("Measured results: ", {"bold": True, "color": GREEN}), f"{N} seeded scenarios, 3 planners each; the script and CSV are in the repository, so no figure is typed in by hand."],
], size=10, bullet="•", after=3)

prs.save(OUT)
print("saved", OUT)
